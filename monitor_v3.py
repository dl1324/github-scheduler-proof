import io
import json
import math
import os
import re
import smtplib
import ssl
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

CONFIG_PATH = Path("config.json")
STATE_PATH = Path("state.json")
TIMEOUT = 35
UA = "uk-mortgage-rate-monitor/3.0 (+GitHub Actions)"


@dataclass
class Obs:
    value: float
    period: str
    source: str


def get(url, timeout=TIMEOUT):
    r = requests.get(url, timeout=timeout, headers={"User-Agent": UA})
    r.raise_for_status()
    return r


def money_number(text):
    m = re.search(r"£\s*([\d,]+(?:\.\d+)?)", text or "")
    return float(m.group(1).replace(",", "")) if m else None


def percent_number(text):
    m = re.search(r"(-?\d+(?:\.\d+)?)\s*%", text or "")
    return float(m.group(1)) if m else None


def latest_ons(path):
    data = get("https://www.ons.gov.uk" + path + "/data").json()
    for bucket in ("months", "quarters", "years"):
        rows = data.get(bucket) or []
        for row in reversed(rows):
            try:
                value = float(str(row["value"]).replace(",", ""))
                period = row.get("date") or row.get("label") or ""
                return Obs(value, period, "ONS")
            except Exception:
                pass
    raise RuntimeError(f"No numeric observation found for {path}")


def bank_rate(fallback):
    try:
        html = get("https://www.bankofengland.co.uk/boeapps/database/Bank-Rate.asp").text
        text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
        for pattern in (
            r"(?:current\s+)?Bank\s+Rate.{0,80}?(\d+(?:\.\d+)?)\s*%",
            r"(\d+(?:\.\d+)?)\s*%[^%]{0,80}Bank\s+Rate",
        ):
            m = re.search(pattern, text, re.I)
            if m:
                return Obs(float(m.group(1)), datetime.now(timezone.utc).date().isoformat(), "Bank of England")
    except Exception:
        pass
    return Obs(float(fallback), "fallback", "config fallback")


def swaps(fallback2, fallback5):
    html = get("https://www.bluegamma.io/sonia-swap-rates-uk").text
    text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
    out = {}
    for tenor, fallback in (("2 Year", fallback2), ("5 Year", fallback5)):
        m = None
        for pattern in (
            rf"{re.escape(tenor)}\s+(?:Live\s+)?(?:\|\s*)?(\d+(?:\.\d+)?)\s*%",
            rf"{re.escape(tenor)}.{{0,90}}?(\d+(?:\.\d+)?)\s*%",
        ):
            m = re.search(pattern, text, re.I)
            if m:
                break
        out[tenor] = Obs(
            float(m.group(1)) if m else float(fallback),
            "previous UK business-day close",
            "BlueGamma public SONIA swap table" if m else "config fallback",
        )
    return out


def _excel_date(value):
    if hasattr(value, "date"):
        try:
            return value.date().isoformat()
        except Exception:
            pass
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            pass
    return str(value)[:10] if value is not None else ""


def boe_gilt_yields(cfg):
    """Read 2Y/5Y/10Y nominal gilt spot yields from the BoE latest yield-curve ZIP."""
    fallback = {y: float(cfg["baseline"][f"gilt_{y}y"]) for y in (2, 5, 10)}
    url = "https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/latest-yield-curve-data.zip"
    try:
        r = get(url, timeout=90)
        z = zipfile.ZipFile(io.BytesIO(r.content))
        names = z.namelist()
        candidates = [
            n for n in names
            if "GLC Nominal" in n and "daily" in n.lower() and n.lower().endswith(".xlsx")
        ]
        if not candidates:
            raise RuntimeError("BoE nominal daily workbook not found")
        wb = load_workbook(io.BytesIO(z.read(candidates[0])), read_only=True, data_only=True)
        sheet_name = next((s for s in wb.sheetnames if s.strip().lower().endswith("spot curve")), None)
        if not sheet_name:
            raise RuntimeError("BoE spot curve sheet not found")
        ws = wb[sheet_name]

        maturity_row = None
        maturity_values = None
        for row_idx in (4, 3, 5, 6):
            vals = [ws.cell(row_idx, c).value for c in range(2, ws.max_column + 1)]
            numeric = []
            for v in vals:
                try:
                    numeric.append(float(v))
                except Exception:
                    numeric.append(None)
            good = [v for v in numeric if v is not None and v > 0]
            if len(good) >= 5:
                maturity_row = row_idx
                maturity_values = numeric
                break
        if maturity_row is None:
            raise RuntimeError("BoE maturity row not detected")

        col_for = {}
        for target in (2.0, 5.0, 10.0):
            matches = [
                (abs(v - target), idx + 2)
                for idx, v in enumerate(maturity_values)
                if v is not None
            ]
            if not matches:
                raise RuntimeError(f"No maturity near {target} years")
            dist, col = min(matches)
            if dist > 0.26:
                raise RuntimeError(f"No exact-enough maturity for {target} years")
            col_for[int(target)] = col

        latest = None
        data_start = maturity_row + 2
        for row in range(data_start, ws.max_row + 1):
            date_val = ws.cell(row, 1).value
            vals = {}
            valid = True
            for years, col in col_for.items():
                try:
                    vals[years] = float(ws.cell(row, col).value)
                except Exception:
                    valid = False
                    break
            if valid:
                latest = (_excel_date(date_val), vals)
        if latest is None:
            raise RuntimeError("No valid BoE gilt observation found")

        period, vals = latest
        return {
            f"gilt_{y}y": Obs(vals[y], period, "Bank of England nominal gilt spot curve")
            for y in (2, 5, 10)
        }
    except Exception:
        return {
            f"gilt_{y}y": Obs(fallback[y], "fallback", "config fallback")
            for y in (2, 5, 10)
        }


def rightmove_average_rates(cfg):
    """Average home-buyer fixed rates by LTV: market-health/calibration input."""
    fallback = cfg["mortgage_market_fallback"]
    result = {str(k): {"2y": float(v["2y"]), "5y": float(v["5y"])} for k, v in fallback.items()}
    source = "config fallback"
    period = "fallback"
    try:
        url = "https://www.rightmove.co.uk/news/articles/property-news/current-uk-mortgage-rates/"
        soup = BeautifulSoup(get(url).text, "html.parser")
        text = soup.get_text(" ", strip=True)
        mdate = re.search(r"(?:Updated:|September|October)\s*([A-Za-z]+\s+\d{1,2},\s+\d{4})", text)
        if mdate:
            period = mdate.group(1)
        found = {}
        for tr in soup.find_all("tr"):
            cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
            if len(cells) < 4:
                continue
            m_ltv = re.fullmatch(r"(60|75|90|95)\s*%", cells[0])
            if not m_ltv:
                continue
            term_text = cells[1].lower().replace("–", "-")
            term = "2y" if "2-year" in term_text else "5y" if "5-year" in term_text else None
            if not term:
                continue
            pcts = []
            for cell in cells[2:]:
                pcts += [float(x) for x in re.findall(r"(-?\d+(?:\.\d+)?)\s*%", cell)]
            if len(pcts) >= 2:
                current = pcts[1]
                if 2.0 <= current <= 10.0:
                    found[(m_ltv.group(1), term)] = current
        if len(found) >= 6:
            for (ltv, term), rate in found.items():
                result.setdefault(ltv, {})[term] = rate
            source = "Rightmove / Podium average home-buyer rates"
            if period == "fallback":
                m = re.search(r"Updated:\s*([A-Za-z]+\s+\d{1,2},\s+\d{4})", text)
                if m:
                    period = m.group(1)
    except Exception:
        pass
    return {"rates": result, "period": period, "source": source}


def _which_table_product(soup, heading_phrase, ltv):
    heading = None
    for tag in soup.find_all(["h2", "h3", "h4"]):
        if heading_phrase.lower() in tag.get_text(" ", strip=True).lower():
            heading = tag
            break
    if not heading:
        return None
    table = heading.find_next("table")
    if not table:
        return None
    for tr in table.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
        if len(cells) < 5:
            continue
        if cells[0].strip() != f"{ltv}%":
            continue
        rate = percent_number(cells[3])
        fee = money_number(cells[4])
        if rate is None:
            continue
        return {
            "lender": cells[1],
            "rate": float(rate),
            "fee": float(fee or 0.0),
            "source": "Which? / Moneyfacts daily remortgage table",
        }
    return None


def which_decision_products(cfg):
    """Competitive remortgage products for fair 2Y/5Y/tracker decision comparisons."""
    ltv = int(cfg["decision_assumptions"]["ltv"])
    fb = cfg["decision_product_fallback"]
    out = {
        "fixed_2y": dict(fb["fixed_2y"]),
        "fixed_5y": dict(fb["fixed_5y"]),
        "tracker_cheapest": dict(fb["tracker_cheapest"]),
        "period": "fallback",
        "source": "config fallback",
    }
    try:
        url = "https://www.which.co.uk/money/mortgages-and-property/mortgages/best-mortgage-rates-and-deals-aLbQB2O2lDAz"
        soup = BeautifulSoup(get(url).text, "html.parser")
        p2 = _which_table_product(soup, "best two-year fixed-rate mortgages for remortgaging", ltv)
        p5 = _which_table_product(soup, "best five-year fixed-rate mortgages for remortgaging", ltv)
        pt = _which_table_product(soup, "best two-year tracker mortgages for remortgaging", ltv)
        if p2:
            out["fixed_2y"] = p2
        if p5:
            out["fixed_5y"] = p5
        if pt:
            out["tracker_cheapest"] = pt
        if p2 and p5 and pt:
            out["source"] = "Which? tables using Moneyfacts data"
            out["period"] = "updated daily"
    except Exception:
        pass
    return out


def flexible_tracker(cfg, bank_rate_value):
    """
    Flexible tracker used for 'wait before fixing' scenarios.
    Nationwide trackers have no ERC according to Nationwide's intermediary mortgage features.
    Rate is sourced from the current remortgage tracker table where possible.
    """
    fb = dict(cfg["flexible_tracker_fallback"])
    out = {
        "lender": fb["lender"],
        "rate": float(fb["rate"]),
        "fee": float(fb["fee"]),
        "margin": float(fb.get("margin", float(fb["rate"]) - bank_rate_value)),
        "erc": 0.0,
        "source": "config fallback; no-ERC policy from Nationwide",
        "period": "fallback",
    }
    try:
        url = "https://hoa.org.uk/best-mortgage-rates/"
        soup = BeautifulSoup(get(url).text, "html.parser")
        heading = None
        for tag in soup.find_all(["h2", "h3", "h4"]):
            txt = tag.get_text(" ", strip=True).lower()
            if "best 2 year tracker" in txt and "remortgage" in txt:
                heading = tag
                break
        table = heading.find_next("table") if heading else None
        if table:
            for tr in table.find_all("tr"):
                cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
                if len(cells) < 7:
                    continue
                if "nationwide" not in cells[0].lower():
                    continue
                if "60%" not in cells[6]:
                    continue
                rate = percent_number(cells[1])
                fee = money_number(cells[2])
                if rate is not None:
                    out.update(
                        lender=cells[0],
                        rate=float(rate),
                        fee=float(fee or fb["fee"]),
                        margin=round(float(rate) - bank_rate_value, 4),
                        source="HomeOwners Alliance / MAB tracker table; Nationwide no-ERC policy",
                        period="current table",
                    )
                    break
    except Exception:
        pass
    return out


def band(name, value, cfg):
    t = cfg["triggers"][name]
    if t.get("direction") == "inverse":
        if value >= t["lower_rate"]:
            return "LOWER"
        if value <= t["higher_rate"]:
            return "HIGHER"
    else:
        if value <= t["lower_rate"]:
            return "LOWER"
        if value >= t["higher_rate"]:
            return "HIGHER"
    return "NEUTRAL"


def pressure_index(values, cfg):
    total = 0.0
    denom = 0.0
    bands = {}
    for key, weight in cfg["weights"].items():
        if key not in values:
            continue
        b = band(key, values[key]["value"], cfg)
        bands[key] = b
        total += weight * (-1 if b == "LOWER" else 1 if b == "HIGHER" else 0)
        denom += weight
    return (total / denom if denom else 0.0), bands


def chosen_average_rates(mortgage_market, cfg):
    ltv = str(cfg["decision_assumptions"]["ltv"])
    row = mortgage_market["rates"].get(ltv) or cfg["mortgage_market_fallback"][ltv]
    return float(row["2y"]), float(row["5y"])


def make_forecast(values, mortgage_market, pressure, cfg):
    b = cfg["baseline"]
    d2 = values["swap_2y"]["value"] - b["swap_2y"]
    d5 = values["swap_5y"]["value"] - b["swap_5y"]
    gilt_delta = (
        0.25 * (values["gilt_2y"]["value"] - b["gilt_2y"])
        + 0.35 * (values["gilt_5y"]["value"] - b["gilt_5y"])
        + 0.40 * (values["gilt_10y"]["value"] - b["gilt_10y"])
    )
    current2, current5 = chosen_average_rates(mortgage_market, cfg)
    ltv = str(cfg["decision_assumptions"]["ltv"])
    base_mort = cfg["mortgage_market_fallback"][ltv]
    market_delta2 = current2 - float(base_mort["2y"])
    market_delta5 = current5 - float(base_mort["5y"])

    result = {}
    for horizon, hcfg in cfg["forecast_anchors"].items():
        shift2 = (
            pressure * hcfg["pressure_sensitivity"]
            + d2 * hcfg["swap_sensitivity"]
            + gilt_delta * hcfg["gilt_sensitivity"]
            + market_delta2 * hcfg["current_market_sensitivity"]
        )
        shift5 = (
            pressure * hcfg["pressure_sensitivity"]
            + d5 * hcfg["swap_sensitivity"]
            + gilt_delta * hcfg["gilt_sensitivity"]
            + market_delta5 * hcfg["current_market_sensitivity"]
        )
        mid2 = round(hcfg["mortgage_2y"] + shift2, 2)
        mid5 = round(hcfg["mortgage_5y"] + shift5, 2)
        width = hcfg["range_half_width"]
        result[horizon] = {
            "mortgage_2y_mid": mid2,
            "mortgage_2y_range": [round(mid2 - width, 2), round(mid2 + width, 2)],
            "mortgage_5y_mid": mid5,
            "mortgage_5y_range": [round(mid5 - width, 2), round(mid5 + width, 2)],
        }
    return result


def bank_rate_forecast(values, pressure, cfg):
    b = cfg["baseline"]
    swap_delta = 0.75 * (values["swap_2y"]["value"] - b["swap_2y"]) + 0.25 * (values["swap_5y"]["value"] - b["swap_5y"])
    out = {}
    for horizon, hcfg in cfg["bank_rate_forecast_anchors"].items():
        mid = hcfg["bank_rate"] + pressure * hcfg["pressure_sensitivity"] + swap_delta * hcfg["swap_sensitivity"]
        out[horizon] = round(max(0.0, mid), 2)
    return out


def interp_horizon(months, current, h6, h12, h24):
    points = [(0, current), (6, h6), (12, h12), (24, h24)]
    if months <= 0:
        return current
    if months >= 24:
        return h24
    for (m0, v0), (m1, v1) in zip(points, points[1:]):
        if m0 <= months <= m1:
            frac = (months - m0) / (m1 - m0)
            return v0 + frac * (v1 - v0)
    return h24


def monthly_interest_and_balance(balance, annual_rate, months, remaining_term_months):
    monthly = annual_rate / 100.0 / 12.0
    payment = balance / remaining_term_months if monthly == 0 else balance * monthly / (1 - (1 + monthly) ** (-remaining_term_months))
    interest_total = 0.0
    bal = balance
    for _ in range(months):
        interest = bal * monthly
        principal = payment - interest
        bal = max(0.0, bal - principal)
        interest_total += interest
    return interest_total, bal


def simulate_variable(balance, term_months, annual_rates):
    bal = balance
    interest_total = 0.0
    months_elapsed = 0
    for annual_rate in annual_rates:
        remaining = max(1, term_months - months_elapsed)
        monthly = annual_rate / 100.0 / 12.0
        payment = bal / remaining if monthly == 0 else bal * monthly / (1 - (1 + monthly) ** (-remaining))
        interest = bal * monthly
        principal = payment - interest
        bal = max(0.0, bal - principal)
        interest_total += interest
        months_elapsed += 1
    return interest_total, bal


def fixed_forecast_for_month(months, term, forecast, current_average, current_best):
    key = "mortgage_2y_mid" if term == "2y" else "mortgage_5y_mid"
    future_avg = interp_horizon(
        months,
        current_average,
        forecast["6 months"][key],
        forecast["12 months"][key],
        forecast["2 years"][key],
    )
    spread = current_best - current_average
    return round(future_avg + spread, 2)


def tracker_monthly_path(wait_months, bank_now, bank_fc, margin):
    rates = []
    for m in range(1, wait_months + 1):
        br = interp_horizon(
            m,
            bank_now,
            bank_fc["6 months"],
            bank_fc["12 months"],
            bank_fc["2 years"],
        )
        rates.append(max(0.0, br + margin))
    return rates


def cost_five_year_fix_now(balance, term_months, rate, fee):
    interest, _ = monthly_interest_and_balance(balance, rate, 60, term_months)
    return interest + fee


def cost_two_year_strategy(balance, term_months, rate2, fee2, refi_rate, refi_fee):
    interest2, bal2 = monthly_interest_and_balance(balance, rate2, 24, term_months)
    interest3, _ = monthly_interest_and_balance(bal2, refi_rate, 36, term_months - 24)
    return interest2 + interest3 + fee2 + refi_fee


def tracker_then_fix_cost(balance, term_months, wait_months, tracker_rates, tracker_fee, future_5y_rate, future_fix_fee, exit_cost=0.0):
    tracker_interest, bal = simulate_variable(balance, term_months, tracker_rates)
    remaining_eval = 60 - wait_months
    fixed_interest, _ = monthly_interest_and_balance(bal, future_5y_rate, remaining_eval, term_months - wait_months)
    return tracker_interest + fixed_interest + tracker_fee + future_fix_fee + exit_cost


def tracker_break_even_future_fix(balance, term_months, wait_months, tracker_rates, tracker_fee, future_fix_fee, exit_cost, cost_5y_now):
    tracker_interest, bal = simulate_variable(balance, term_months, tracker_rates)
    remaining_eval = 60 - wait_months

    def total(rate):
        fixed_interest, _ = monthly_interest_and_balance(bal, rate, remaining_eval, term_months - wait_months)
        return tracker_interest + fixed_interest + tracker_fee + future_fix_fee + exit_cost

    lo, hi = 0.0, 15.0
    for _ in range(60):
        mid = (lo + hi) / 2.0
        if total(mid) < cost_5y_now:
            lo = mid
        else:
            hi = mid
    return round((lo + hi) / 2.0, 2)


def decision_view(mortgage_market, decision_products, flex_tracker, forecast, bank_fc, bank_now, cfg):
    a = cfg["decision_assumptions"]
    balance = float(a["loan_amount"])
    term_months = int(a["term_years"]) * 12
    ltv = str(a["ltv"])

    avg2, avg5 = chosen_average_rates(mortgage_market, cfg)
    p2 = decision_products["fixed_2y"]
    p5 = decision_products["fixed_5y"]
    current2 = float(p2["rate"])
    current5 = float(p5["rate"])
    fee2 = float(p2["fee"])
    fee5 = float(p5["fee"])
    refi_fee = float(a["refinance_fee"])

    refi2 = fixed_forecast_for_month(24, "2y", forecast, avg2, current2)
    refi5 = fixed_forecast_for_month(24, "5y", forecast, avg5, current5)
    refi_proxy = round((refi2 + refi5) / 2.0, 2)

    cost5 = cost_five_year_fix_now(balance, term_months, current5, fee5)
    cost2 = cost_two_year_strategy(balance, term_months, current2, fee2, refi_proxy, refi_fee)

    waits = []
    margin = float(flex_tracker["margin"])
    for wait in a["tracker_wait_months"]:
        tracker_rates = tracker_monthly_path(int(wait), bank_now, bank_fc, margin)
        future_5 = fixed_forecast_for_month(int(wait), "5y", forecast, avg5, current5)
        future_fix_fee = float(a.get("future_fix_fee", fee5))
        wait_cost = tracker_then_fix_cost(
            balance,
            term_months,
            int(wait),
            tracker_rates,
            float(flex_tracker["fee"]),
            future_5,
            future_fix_fee,
            float(a.get("tracker_exit_cost", 0.0)),
        )
        be = tracker_break_even_future_fix(
            balance,
            term_months,
            int(wait),
            tracker_rates,
            float(flex_tracker["fee"]),
            future_fix_fee,
            float(a.get("tracker_exit_cost", 0.0)),
            cost5,
        )
        if future_5 <= be - 0.10:
            signal = "Model path is below break-even: waiting on the flexible tracker is cheaper than a 5Y fix now under these assumptions."
        elif future_5 >= be + 0.10:
            signal = "Model path is above break-even: the 5Y fix now is cheaper than waiting under these assumptions."
        else:
            signal = "Close to break-even: fees, product eligibility and small rate moves can change the result."
        waits.append({
            "months": int(wait),
            "projected_tracker_rate_at_fix": round(tracker_rates[-1], 2),
            "model_5y_fix_at_wait": future_5,
            "break_even_5y_fix_at_wait": be,
            "expected_5y_cost": round(wait_cost, 0),
            "signal": signal,
        })

    strategies = [
        {"name": "5Y fix now", "expected_cost": round(cost5, 0)},
        {"name": "2Y fix now + modelled refinance", "expected_cost": round(cost2, 0)},
    ] + [
        {"name": f"Flexible tracker {w['months']}m → 5Y fix", "expected_cost": w["expected_5y_cost"]}
        for w in waits
    ]
    best_cost = min(x["expected_cost"] for x in strategies)
    for x in strategies:
        x["difference_vs_lowest"] = round(x["expected_cost"] - best_cost, 0)

    def cost2_then(rate):
        return cost_two_year_strategy(balance, term_months, current2, fee2, rate, refi_fee)

    lo, hi = 0.0, 15.0
    for _ in range(60):
        mid = (lo + hi) / 2.0
        if cost2_then(mid) < cost5:
            lo = mid
        else:
            hi = mid
    be_2v5 = round((lo + hi) / 2.0, 2)

    return {
        "ltv": int(ltv),
        "fixed_2y": p2,
        "fixed_5y": p5,
        "tracker_cheapest": decision_products["tracker_cheapest"],
        "tracker_flexible": flex_tracker,
        "break_even_refi_rate_2v5": be_2v5,
        "model_refi_proxy_2y": refi_proxy,
        "wait_scenarios": waits,
        "strategies": strategies,
    }


def fmt(v):
    return f"{float(v):.2f}%"


def fmt_money(v):
    return f"£{float(v):,.0f}"


def changed_enough(a, b, threshold):
    try:
        return abs(float(a) - float(b)) >= threshold
    except Exception:
        return True


def reasons(values, bands, mortgage_market, decision_products, flex_tracker, prev, forecast, bank_fc, cfg):
    out = []
    anchor = (prev or {}).get("alert_anchor", {}) or {}
    old_bands = anchor.get("bands", {})
    old_vals = anchor.get("values", {})
    old_fc = anchor.get("forecast", {})
    old_bank_fc = anchor.get("bank_rate_forecast", {})
    old_market = anchor.get("mortgage_market", {})
    old_products = anchor.get("decision_products", {})
    old_flex = anchor.get("flexible_tracker", {})

    for key, b in bands.items():
        if old_bands and key not in old_bands:
            out.append(f"New trigger added: {cfg['labels'].get(key, key)} = {b}")
        elif old_bands and old_bands.get(key) != b:
            out.append(f"{cfg['labels'].get(key, key)} trigger moved {old_bands.get(key, '?')} → {b}")

    for key in ("swap_2y", "swap_5y"):
        if key in old_vals and changed_enough(values[key]["value"], old_vals[key]["value"], cfg["alert_rules"]["swap_move_pp"]):
            out.append(f"{cfg['labels'][key]} moved {values[key]['value'] - old_vals[key]['value']:+.2f}pp since last alert")

    for key in ("gilt_2y", "gilt_5y", "gilt_10y"):
        if key in old_vals and changed_enough(values[key]["value"], old_vals[key]["value"], cfg["alert_rules"]["gilt_move_pp"]):
            out.append(f"{cfg['labels'][key]} moved {values[key]['value'] - old_vals[key]['value']:+.2f}pp since last alert")

    if "bank_rate" in old_vals and changed_enough(values["bank_rate"]["value"], old_vals["bank_rate"]["value"], 0.24):
        out.append(f"Bank Rate changed from {old_vals['bank_rate']['value']:.2f}% to {values['bank_rate']['value']:.2f}%")

    ltv = str(cfg["decision_assumptions"]["ltv"])
    try:
        old_rates = old_market["rates"][ltv]
        new_rates = mortgage_market["rates"][ltv]
        for term in ("2y", "5y"):
            if changed_enough(new_rates[term], old_rates[term], cfg["alert_rules"]["mortgage_move_pp"]):
                out.append(f"{ltv}% LTV average {term.upper()} mortgage moved {new_rates[term] - old_rates[term]:+.2f}pp since last alert")
    except Exception:
        pass

    for product_key, label in (("fixed_2y", "competitive 2Y fix"), ("fixed_5y", "competitive 5Y fix"), ("tracker_cheapest", "cheapest tracker")):
        try:
            old_rate = old_products[product_key]["rate"]
            new_rate = decision_products[product_key]["rate"]
            if changed_enough(new_rate, old_rate, cfg["alert_rules"]["mortgage_move_pp"]):
                out.append(f"{label} moved {new_rate - old_rate:+.2f}pp since last alert")
        except Exception:
            pass

    try:
        if changed_enough(flex_tracker["rate"], old_flex["rate"], cfg["alert_rules"]["mortgage_move_pp"]):
            out.append(f"Flexible tracker proxy moved {flex_tracker['rate'] - old_flex['rate']:+.2f}pp since last alert")
    except Exception:
        pass

    for key, value in values.items():
        if "fallback" in value.get("source", "").lower():
            old_source = old_vals.get(key, {}).get("source", "")
            if "fallback" not in old_source.lower():
                out.append(f"Data-source warning: {cfg['labels'].get(key, key)} is using its configured fallback")

    for horizon, row in forecast.items():
        if horizon in old_fc:
            for term in ("mortgage_2y_mid", "mortgage_5y_mid"):
                if changed_enough(row[term], old_fc[horizon][term], cfg["alert_rules"]["forecast_move_pp"]):
                    out.append(f"{horizon} fixed-mortgage forecast moved materially")
                    break

    for horizon, value in bank_fc.items():
        if horizon in old_bank_fc and changed_enough(value, old_bank_fc[horizon], cfg["alert_rules"]["forecast_move_pp"]):
            out.append(f"{horizon} Bank Rate model moved materially")

    return list(dict.fromkeys(out))


def outlook_direction(current2, current5, forecast):
    current = (current2 + current5) / 2.0
    six = (forecast["6 months"]["mortgage_2y_mid"] + forecast["6 months"]["mortgage_5y_mid"]) / 2.0
    delta = six - current
    if delta <= -0.50:
        return "Falling rapidly"
    if delta <= -0.10:
        return "Falling gradually"
    if delta < 0.10:
        return "Broadly stable"
    if delta < 0.50:
        return "Rising gradually"
    return "Rising rapidly"


def email_html(values, bands, mortgage_market, forecast, bank_fc, decision, why, pressure, cfg):
    ordered = [
        ("bank_rate", "Bank Rate"),
        ("headline_cpi", "Headline CPI"),
        ("core_cpi", "Core CPI"),
        ("services_cpi", "Services CPI"),
        ("wage_growth", "Private regular wage growth"),
        ("unemployment", "Unemployment"),
        ("swap_2y", "2Y SONIA swap"),
        ("swap_5y", "5Y SONIA swap"),
        ("gilt_2y", "2Y gilt yield"),
        ("gilt_5y", "5Y gilt yield"),
        ("gilt_10y", "10Y gilt yield"),
    ]
    rows = ""
    for key, label in ordered:
        v = values.get(key)
        if v:
            rows += f"<tr><td>{label}</td><td>{fmt(v['value'])}</td><td>{bands.get(key, 'TRACK')}</td><td>{v['period']}</td><td>{v['source']}</td></tr>"

    market_rows = ""
    for ltv in ("60", "75", "90"):
        if ltv in mortgage_market["rates"]:
            x = mortgage_market["rates"][ltv]
            market_rows += f"<tr><td>{ltv}%</td><td>{fmt(x['2y'])}</td><td>{fmt(x['5y'])}</td></tr>"

    frows = ""
    for h, x in forecast.items():
        tracker_mid = bank_fc[h] + float(decision["tracker_flexible"]["margin"])
        frows += (
            f"<tr><td>{h}</td><td>{fmt(bank_fc[h])}</td>"
            f"<td>{fmt(x['mortgage_2y_mid'])} ({fmt(x['mortgage_2y_range'][0])}–{fmt(x['mortgage_2y_range'][1])})</td>"
            f"<td>{fmt(x['mortgage_5y_mid'])} ({fmt(x['mortgage_5y_range'][0])}–{fmt(x['mortgage_5y_range'][1])})</td>"
            f"<td>{fmt(tracker_mid)}</td></tr>"
        )

    strategy_rows = "".join(
        f"<tr><td>{x['name']}</td><td>{fmt_money(x['expected_cost'])}</td><td>{fmt_money(x['difference_vs_lowest'])}</td></tr>"
        for x in decision["strategies"]
    )
    wait_rows = "".join(
        f"<tr><td>{w['months']} months</td><td>{fmt(w['projected_tracker_rate_at_fix'])}</td>"
        f"<td>{fmt(w['model_5y_fix_at_wait'])}</td><td>{fmt(w['break_even_5y_fix_at_wait'])}</td><td>{w['signal']}</td></tr>"
        for w in decision["wait_scenarios"]
    )
    why_html = "".join(f"<li>{r}</li>" for r in why) or "<li>Manual/test report</li>"
    signal = "higher-rate pressure" if pressure > 0.15 else "lower-rate pressure" if pressure < -0.15 else "broadly neutral"
    direction = outlook_direction(decision["fixed_2y"]["rate"], decision["fixed_5y"]["rate"], forecast)
    a = cfg["decision_assumptions"]

    return f"""<html><body>
    <h2>UK Mortgage Outlook & Decision Monitor</h2>
    <p><b>Current direction:</b> {direction}<br><b>Macro/market pressure:</b> {signal} ({pressure:+.2f})</p>

    <h3>Current mortgage market averages</h3>
    <table border="1" cellpadding="6" cellspacing="0"><tr><th>LTV</th><th>Average 2Y fix</th><th>Average 5Y fix</th></tr>{market_rows}</table>
    <p><small>{mortgage_market['source']} — {mortgage_market['period']}. Market-health/calibration data, not the products used in the strategy comparison.</small></p>

    <h3>Competitive {decision['ltv']}% LTV remortgage products used for decisions</h3>
    <table border="1" cellpadding="6" cellspacing="0">
      <tr><th>Route</th><th>Rate</th><th>Fee</th><th>Lender / note</th></tr>
      <tr><td>2Y fix</td><td>{fmt(decision['fixed_2y']['rate'])}</td><td>{fmt_money(decision['fixed_2y']['fee'])}</td><td>{decision['fixed_2y']['lender']}</td></tr>
      <tr><td>5Y fix</td><td>{fmt(decision['fixed_5y']['rate'])}</td><td>{fmt_money(decision['fixed_5y']['fee'])}</td><td>{decision['fixed_5y']['lender']}</td></tr>
      <tr><td>Cheapest tracker</td><td>{fmt(decision['tracker_cheapest']['rate'])}</td><td>{fmt_money(decision['tracker_cheapest']['fee'])}</td><td>{decision['tracker_cheapest']['lender']} — exit terms must be checked</td></tr>
      <tr><td>Flexible tracker proxy</td><td>{fmt(decision['tracker_flexible']['rate'])}</td><td>{fmt_money(decision['tracker_flexible']['fee'])}</td><td>{decision['tracker_flexible']['lender']} — model assumes no ERC</td></tr>
    </table>

    <h3>Forecast</h3>
    <table border="1" cellpadding="6" cellspacing="0"><tr><th>Horizon</th><th>Bank Rate model</th><th>Typical 2Y fix</th><th>Typical 5Y fix</th><th>Flexible tracker proxy</th></tr>{frows}</table>

    <h3>5-year cost comparison</h3>
    <p>Interest + modelled product/refinance fees over the next 60 months; principal repayment is not treated as a cost.</p>
    <table border="1" cellpadding="6" cellspacing="0"><tr><th>Strategy</th><th>Expected interest + fees</th><th>Above lowest model cost</th></tr>{strategy_rows}</table>

    <h3>Should a flexible tracker be used to wait before fixing?</h3>
    <table border="1" cellpadding="6" cellspacing="0"><tr><th>Wait</th><th>Projected tracker rate at switch</th><th>Modelled 5Y fix then</th><th>Break-even 5Y fix</th><th>Model signal</th></tr>{wait_rows}</table>
    <p><small>The wait strategy is only practical if the actual tracker permits switching/repayment at the chosen time without a material ERC. The model deliberately uses a flexible/no-ERC tracker proxy rather than blindly using the absolute cheapest tracker.</small></p>

    <h3>2Y vs 5Y fix now</h3>
    <p>For the 2Y-fix route to beat the 5Y fix over five years, the modelled refinance rate in two years needs to be below about <b>{fmt(decision['break_even_refi_rate_2v5'])}</b>. Current model proxy: <b>{fmt(decision['model_refi_proxy_2y'])}</b>.</p>

    <p><small>Decision assumptions: £{a['loan_amount']:,} repayment mortgage, {a['term_years']}-year remaining term, {a['ltv']}% LTV, refinance fee £{a['refinance_fee']:,}, future fix fee £{a['future_fix_fee']:,}, tracker exit cost £{a['tracker_exit_cost']:,}. Product-specific legal/valuation/broker costs are not automatically included.</small></p>

    <h3>What changed</h3><ul>{why_html}</ul>
    <h3>Market & economic dashboard</h3>
    <table border="1" cellpadding="6" cellspacing="0"><tr><th>Indicator</th><th>Latest</th><th>Band</th><th>Period</th><th>Source</th></tr>{rows}</table>
    <p><small>LOWER = signal consistent with lower mortgage rates; HIGHER = signal consistent with higher rates. SONIA swaps are the primary fixed-mortgage market signal. BoE gilt yields add term-premium/fiscal/global-bond information. Forecasts are scenario-model ranges, not guaranteed lender quotes or regulated mortgage advice.</small></p>
    </body></html>"""


def send_email(subject, html):
    email_from = os.environ.get("EMAIL_FROM")
    email_to = os.environ.get("EMAIL_TO")
    password = os.environ.get("EMAIL_APP_PASSWORD")
    if not all((email_from, email_to, password)):
        raise RuntimeError("Missing EMAIL_FROM, EMAIL_TO or EMAIL_APP_PASSWORD GitHub secret")
    msg = EmailMessage()
    msg["From"] = email_from
    msg["To"] = email_to
    msg["Subject"] = subject
    msg.set_content("UK Mortgage Outlook & Decision Monitor - view this email in HTML.")
    msg.add_alternative(html, subtype="html")
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as smtp:
        smtp.login(email_from, password)
        smtp.send_message(msg)


def main():
    cfg = json.loads(CONFIG_PATH.read_text())
    prev = json.loads(STATE_PATH.read_text()) if STATE_PATH.exists() else {}
    values = {}

    series = {
        "headline_cpi": "/economy/inflationandpriceindices/timeseries/d7g7/mm23",
        "core_cpi": "/economy/inflationandpriceindices/timeseries/dko8/mm23",
        "services_cpi": "/economy/inflationandpriceindices/timeseries/d7nn/mm23",
        "wage_growth": "/employmentandlabourmarket/peopleinwork/earningsandworkinghours/timeseries/kaj4/lms",
        "unemployment": "/employmentandlabourmarket/peoplenotinwork/unemployment/timeseries/mgsx/lms",
    }
    for key, path in series.items():
        try:
            obs = latest_ons(path)
        except Exception:
            obs = Obs(cfg["baseline"][key], "fallback", "config fallback")
        values[key] = obs.__dict__

    values["bank_rate"] = bank_rate(cfg["baseline"]["bank_rate"]).__dict__

    try:
        swap_obs = swaps(cfg["baseline"]["swap_2y"], cfg["baseline"]["swap_5y"])
        values["swap_2y"] = swap_obs["2 Year"].__dict__
        values["swap_5y"] = swap_obs["5 Year"].__dict__
    except Exception:
        values["swap_2y"] = Obs(cfg["baseline"]["swap_2y"], "fallback", "config fallback").__dict__
        values["swap_5y"] = Obs(cfg["baseline"]["swap_5y"], "fallback", "config fallback").__dict__

    for key, obs in boe_gilt_yields(cfg).items():
        values[key] = obs.__dict__

    mortgage_market = rightmove_average_rates(cfg)
    decision_products = which_decision_products(cfg)
    flex = flexible_tracker(cfg, values["bank_rate"]["value"])

    pressure, bands = pressure_index(values, cfg)
    forecast = make_forecast(values, mortgage_market, pressure, cfg)
    bank_fc = bank_rate_forecast(values, pressure, cfg)
    decision = decision_view(
        mortgage_market,
        decision_products,
        flex,
        forecast,
        bank_fc,
        values["bank_rate"]["value"],
        cfg,
    )
    why = reasons(values, bands, mortgage_market, decision_products, flex, prev, forecast, bank_fc, cfg)

    test = os.environ.get("TEST_EMAIL", "").lower() in {"1", "true", "yes"}
    first = not prev.get("alert_anchor")
    should_send = test or first or bool(why)
    now = datetime.now(timezone.utc).isoformat()

    state = {
        "last_checked": now,
        "latest_values": values,
        "mortgage_market": mortgage_market,
        "decision_products": decision_products,
        "flexible_tracker": flex,
        "latest_bands": bands,
        "latest_pressure_index": round(pressure, 4),
        "latest_forecast": forecast,
        "latest_bank_rate_forecast": bank_fc,
        "latest_decision_view": decision,
        "alert_anchor": prev.get("alert_anchor"),
    }

    if should_send:
        prefix = "TEST — " if test else ""
        subject = f"{prefix}UK Mortgage Outlook: {outlook_direction(decision['fixed_2y']['rate'], decision['fixed_5y']['rate'], forecast)}"
        send_email(subject, email_html(values, bands, mortgage_market, forecast, bank_fc, decision, why, pressure, cfg))
        state["alert_anchor"] = {
            "values": values,
            "mortgage_market": mortgage_market,
            "decision_products": decision_products,
            "flexible_tracker": flex,
            "bands": bands,
            "forecast": forecast,
            "bank_rate_forecast": bank_fc,
            "decision_view": decision,
            "sent_at": now,
        }

    STATE_PATH.write_text(json.dumps(state, indent=2) + "\n")


if __name__ == "__main__":
    main()
